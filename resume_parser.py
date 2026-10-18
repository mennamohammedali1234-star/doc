"""
Resume Parser and NLP Analysis Utilities
Extracts text from PDF, DOCX, and TXT files, and analyzes keyword matches
including negation detection and competency mapping.
"""

import io
import re
from typing import BinaryIO, Union
from main import (
    detect_negations,
    extract_competencies,
    normalize_text,
    stem_phrase,
    stemmer,
    match_competency_in_resume,
    TECH_ALIASES
)


def extract_text_from_file(file_source: Union[str, BinaryIO], filename: str) -> str:
    """Extract text from PDF, DOCX, or plain text file."""
    filename_lower = filename.lower()

    if filename_lower.endswith(".pdf"):
        import pypdf

        if isinstance(file_source, str):
            reader = pypdf.PdfReader(file_source)
        else:
            reader = pypdf.PdfReader(file_source)

        text_pages = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text_pages.append(extracted)
        return "\n".join(text_pages).strip()

    elif filename_lower.endswith(".docx"):
        import docx

        doc = docx.Document(file_source)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs).strip()

    else:
        # Plain text / Markdown / unknown
        if isinstance(file_source, str):
            with open(file_source, "rb") as f:
                content = f.read()
        else:
            content = file_source.read()

        for encoding in ["utf-8", "utf-16", "cp1252", "latin-1"]:
            try:
                return content.decode(encoding).strip()
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore").strip()


def analyze_keyword_matches(job_desc: str, resume_text: str, top_n: int = 15) -> dict:
    """
    Find matched, missing, and explicitly negated competencies between job description
    and candidate resume using morphological stemming and technology thesaurus.
    """
    job_competencies = extract_competencies(job_desc)
    resume_norm = normalize_text(resume_text)
    resume_stemmed_norm = stem_phrase(resume_text)

    cand_tokens = set(re.findall(r"\b[a-zA-Z][a-zA-Z0-9+#.-]{1,25}\b", resume_norm))
    cand_stemmed_tokens = {stemmer.stem(t) for t in cand_tokens}
    negated_set = detect_negations(resume_text)

    matched = []
    missing = []
    negated = []

    for comp in job_competencies:
        comp_stemmed = stem_phrase(comp)
        is_negated = (
            comp in negated_set or
            comp_stemmed in negated_set or
            any(neg in comp for neg in negated_set if len(neg) > 2)
        )

        if is_negated:
            negated.append(comp)
        elif match_competency_in_resume(comp, resume_norm, resume_stemmed_norm, cand_tokens, cand_stemmed_tokens):
            matched.append(comp)
        else:
            missing.append(comp)

    total = max(1, len(job_competencies))
    skill_coverage = len(matched) / total

    return {
        "matched": matched[:top_n],
        "missing": missing[:top_n],
        "negated": negated[:top_n],
        "total_job_competencies": len(job_competencies),
        "matched_count": len(matched),
        "negated_count": len(negated),
        "skill_score_percent": f"{skill_coverage * 100:.1f}%",
        "skills_ratio": f"{len(matched)} of {len(job_competencies)} skills found"
    }
